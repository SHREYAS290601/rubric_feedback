from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".docx", ".txt", ".md", ".pdf"}
SUPPORTED_TYPES_LABEL = ".txt, .md, .docx, and text-based .pdf files"


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type. Supported uploads are {SUPPORTED_TYPES_LABEL}.")

    if extension in {".txt", ".md"}:
        return _normalize_text(path.read_text(encoding="utf-8"))

    if extension == ".pdf":
        return _extract_pdf_text(path)

    return _extract_docx_text(path)


def _extract_docx_text(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []

    blocks.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return _normalize_text("\n\n".join(blocks))


def _extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = _normalize_text("\n\n".join(page for page in pages if page.strip()))
    if not text:
        raise ValueError("This PDF does not contain extractable text. Please upload a text-based PDF or paste the content.")
    return text


def _normalize_text(text: str) -> str:
    normalized = "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"))
    return normalized.strip()
